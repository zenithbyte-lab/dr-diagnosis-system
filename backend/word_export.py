"""Generate Word diagnosis report using python-docx."""
import os
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from .config import REPORT_DIR
from .constants import RISK_CN, GENDER_CN, EYE_CN, SEVERITY_CN, REVIEW_CN

RISK_COLORS = {
    "Low": RGBColor(0x10, 0xB9, 0x81),
    "Low-Medium": RGBColor(0x3B, 0x82, 0xF6),
    "Medium": RGBColor(0xF5, 0x9E, 0x0B),
    "High": RGBColor(0xEF, 0x44, 0x44),
    "Critical": RGBColor(0xDC, 0x26, 0x26),
}


def generate_report(diagnosis) -> str:
    doc = Document()

    # --- Title ---
    title = doc.add_heading("糖尿病视网膜病变诊断报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"报告编号：{diagnosis.uid}    生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    doc.add_paragraph("")

    # --- Patient Information ---
    doc.add_heading("一、患者信息", level=1)
    table = doc.add_table(rows=4, cols=4)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cells = table.rows[0].cells
    cells[0].text = "患者姓名"
    cells[1].text = diagnosis.patient_name or "未填写"
    cells[2].text = "年龄"
    cells[3].text = str(diagnosis.patient_age) if diagnosis.patient_age else "未填写"
    cells = table.rows[1].cells
    cells[0].text = "性别"
    cells[1].text = GENDER_CN.get(diagnosis.patient_gender or "", "未填写")
    cells[2].text = "眼别"
    cells[3].text = EYE_CN.get(diagnosis.eye_side or "", "未指定")
    cells = table.rows[2].cells
    cells[0].text = "诊断时间"
    cells[1].text = diagnosis.created_at.strftime("%Y-%m-%d %H:%M:%S") if diagnosis.created_at else ""
    cells[2].text = "报告编号"
    cells[3].text = diagnosis.uid
    cells = table.rows[3].cells
    cells[0].text = "诊断模型"
    cells[1].text = "ResNet-152（迁移学习）"
    cells[2].text = "模型版本"
    cells[3].text = "v1.0"

    doc.add_paragraph("")

    # --- Fundus Image ---
    doc.add_heading("二、眼底图像", level=1)
    try:
        doc.add_picture(diagnosis.original_image, width=Inches(3.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    except (FileNotFoundError, OSError):
        doc.add_paragraph("[图像不可用]")

    # --- Grad-CAM Heatmap ---
    doc.add_heading("三、Grad-CAM 注意力热力图", level=1)
    doc.add_paragraph("热力图展示了模型在进行诊断判断时重点关注的区域。红色/黄色区域表示模型高度关注的位置（疑似病灶），蓝色/绿色区域表示关注度较低的位置。")
    try:
        doc.add_picture(diagnosis.heatmap_image, width=Inches(3.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    except (FileNotFoundError, OSError):
        doc.add_paragraph("[热力图不可用]")

    # --- Diagnosis Result ---
    doc.add_heading("四、诊断结果", level=1)
    result_table = doc.add_table(rows=3, cols=2)
    result_table.style = "Light Grid Accent 1"
    result_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    result_table.rows[0].cells[0].text = "病变等级"
    pred_idx = int(diagnosis.pred_class)
    run = result_table.rows[0].cells[1].paragraphs[0].add_run(SEVERITY_CN[pred_idx])
    run.bold = True
    color = RISK_COLORS.get(diagnosis.risk_level, RGBColor(0, 0, 0))
    run.font.color.rgb = color

    result_table.rows[1].cells[0].text = "诊断置信度"
    result_table.rows[1].cells[1].text = f"{diagnosis.confidence}%"

    result_table.rows[2].cells[0].text = "风险等级"
    run2 = result_table.rows[2].cells[1].paragraphs[0].add_run(RISK_CN.get(diagnosis.risk_level, diagnosis.risk_level))
    run2.bold = True
    run2.font.color.rgb = color

    doc.add_paragraph("")

    # --- Class Probabilities ---
    doc.add_heading("五、各等级分类概率", level=1)
    prob_table = doc.add_table(rows=6, cols=2)
    prob_table.style = "Light Grid Accent 1"
    prob_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    prob_table.rows[0].cells[0].text = "病变等级"
    prob_table.rows[0].cells[1].text = "概率"
    probs = [
        ("无明显病变（0级）", diagnosis.prob_no_dr),
        ("轻度（1级）", diagnosis.prob_mild),
        ("中度（2级）", diagnosis.prob_moderate),
        ("重度（3级）", diagnosis.prob_severe),
        ("增殖期（4级）", diagnosis.prob_proliferative),
    ]
    for i, (label, prob) in enumerate(probs):
        prob_table.rows[i + 1].cells[0].text = label
        prob_table.rows[i + 1].cells[1].text = f"{prob}%"

    doc.add_paragraph("")

    # --- Clinical Recommendation ---
    doc.add_heading("六、临床建议", level=1)
    doc.add_paragraph(diagnosis.recommendation or "无")

    # --- Review Status ---
    doc.add_heading("七、审核信息", level=1)
    review_status = getattr(diagnosis, 'review_status', 'pending') or 'pending'
    review_table = doc.add_table(rows=2, cols=2)
    review_table.style = "Light Grid Accent 1"
    review_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    review_table.rows[0].cells[0].text = "审核状态"
    review_table.rows[0].cells[1].text = REVIEW_CN.get(review_status, review_status)
    review_table.rows[1].cells[0].text = "审核意见"
    review_table.rows[1].cells[1].text = getattr(diagnosis, 'review_comment', '') or "暂无"

    # --- Disclaimer ---
    doc.add_paragraph("")
    disclaimer = doc.add_paragraph()
    run = disclaimer.add_run(
        "免责声明：本报告由基于深度学习的糖尿病视网膜病变辅助诊断系统自动生成，"
        "仅供临床参考，不构成正式医学诊断。最终诊断结论请以执业眼科医师意见为准。"
    )
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    path = os.path.join(REPORT_DIR, f"{diagnosis.uid}_report.docx")
    doc.save(path)
    return path
